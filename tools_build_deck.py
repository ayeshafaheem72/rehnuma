"""Rehnuma - the five slides for the UBL panel.

The brief allows at most five slides covering: problem and UX, architecture, security,
efficiency and effectiveness, and measurement. Efficiency, effectiveness and measurement share
slide 4 (the brief's own "Efficient & Effective" requirement joins them), which frees slide 5 for a
table that maps all twelve requirements in the brief to how each is covered and the technology
behind it.

Design brief: read as work by someone who has done this professionally for a decade -
restraint, not decoration.
  - White ground; UBL's own blue as the single accent; greys for everything else.
  - Cambria headlines against Calibri body.
  - Every technical slide leads with one plain-English sentence, then the detail underneath,
    so a non-technical panellist and an engineer both get served.
  - The pictures are real screenshots of the running system (docs/screens), not mock-ups.

Run:  python tools_build_deck.py
"""
import math
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJ = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(PROJ, "Rehnuma_UBL_5_Slides.pptx")
LOGO = os.path.join(PROJ, "static", "ubl-logo.png")
SHOTS = os.path.join(PROJ, "docs", "screens")

NAVY = RGBColor(0x0C, 0x2B, 0x45)
BLUE = RGBColor(0x00, 0x70, 0xB8)
BLUE_L = RGBColor(0x50, 0xA0, 0xD0)
DEEP = RGBColor(0x0F, 0x58, 0x85)
GREY = RGBColor(0x6B, 0x6E, 0x72)
GREY_L = RGBColor(0xE6, 0xE7, 0xE8)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
FAINT = RGBColor(0xF4, 0xF7, 0xFA)
RANI = RGBColor(0xE6, 0x00, 0x6E)
GREEN = RGBColor(0x00, 0x80, 0x5C)

HEAD, BODY = "Cambria", "Calibri"
M = 0.85
CW = 13.333 - 2 * M

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(notes=""):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = PAPER
    r.line.fill.background(); r.shadow.inherit = False
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def txt(s, x, y, w, h, runs, size=13, color=GREY, font=BODY, bold=False, align=PP_ALIGN.LEFT,
        space=5, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (t, o) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", space))
        ls = o.get("spacing", spacing)
        if ls:
            p.line_spacing = ls
        r = p.add_run(); r.text = t
        f = r.font
        f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
        f.italic = o.get("italic", False); f.name = o.get("font", font)
        f.color.rgb = o.get("color", color)
    return tb


def fill(s, x, y, w, h, color=FAINT, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def picture(s, name, x, y, w):
    """A screenshot with a hairline frame, so a dark UI does not bleed into a white slide."""
    path = os.path.join(SHOTS, name)
    from PIL import Image
    iw, ih = Image.open(path).size
    h = w * ih / iw
    fill(s, x - 0.03, y - 0.03, w + 0.06, h + 0.06, PAPER, GREY_L)
    s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    return h


def logo(s, x, y, h=0.4):
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(x), Inches(y), height=Inches(h))


def head(s, kicker, headline, size=34, width=CW):
    txt(s, M, 0.62, width, 0.26, kicker, size=10.5, color=BLUE, bold=True)
    txt(s, M, 0.92, width, 0.8, headline, size=size, color=NAVY, font=HEAD, bold=True, spacing=0.95)


def plain(s, y, line, width=CW, size=16):
    txt(s, M, y, width, 0.85, line, size=size, color=DEEP, font=HEAD, spacing=1.12)


def foot(s, n):
    txt(s, M, 7.02, 9.5, 0.26,
        "Ayesha Faheem  ·  Digital HR & AI Manager assignment  ·  prototype for assessment, "
        "not an official UBL service", size=8.5, color=GREY)
    txt(s, 11.0, 7.02, 1.48, 0.26, f"{n} / 5", size=8.5, color=GREY, align=PP_ALIGN.RIGHT)


def arrow(s, x1, y1, x2, y2, label=None, both=False, color=BLUE, lx=0.0, ly=-0.27, lw=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if label:
        mx, my = (x1 + x2) / 2 + lx, (y1 + y2) / 2 + ly
        lines = list(label) if isinstance(label, (list, tuple)) else str(label).splitlines()
        txt(s, mx - lw / 2, my, lw, 0.4, [(t, {"space": 0}) for t in lines], size=8.5, color=DEEP,
            align=PP_ALIGN.CENTER, space=0)


def box(s, x, y, w, h, title, lines=(), tone="plain", tsize=11.5, lsize=9.5):
    bg, edge, tcol = {"plain": (PAPER, GREY_L, NAVY), "blue": (FAINT, BLUE_L, DEEP),
                      "strong": (BLUE, BLUE, PAPER)}[tone]
    fill(s, x, y, w, h, bg, edge)
    body = [(title, {"size": tsize, "bold": True, "color": tcol, "font": HEAD, "space": 3})]
    lcol = PAPER if tone == "strong" else GREY
    for ln in lines:
        body.append((ln, {"size": lsize, "color": lcol, "space": 1.5}))
    txt(s, x + 0.12, y + 0.09, w - 0.24, h - 0.14, body)


# ══════════════════════════════════════════════ 1 · PROBLEM & EXPERIENCE (UX)
s = slide(
    "PROBLEM AND EXPERIENCE. Every training tool ends with a quiz, and a quiz tells you what someone "
    "remembered this morning, not whether they can use it. Rehnuma never tests anyone. You give it any "
    "material; it turns it into an illustrated story you watch, then a game you play, in English or Urdu, "
    "typed or spoken. From how you talk it works out what you understand and changes how it teaches. "
    "The pictures on this slide are the real running system.")
logo(s, M, 0.1, 0.34)
txt(s, M + 0.95, 0.16, 4.0, 0.3, "REHNUMA  ·  رہنما", size=11, color=NAVY, bold=True)
head(s, "PROBLEM & EXPERIENCE", "Nobody ever learned from a quiz", size=30, width=4.9)
plain(s, 2.05,
      "Rehnuma turns any document into a story you watch, then a game you play, in English or Urdu. "
      "It works out what you understand from how you talk — and never gives a test.",
      width=5.7, size=15.5)

steps = [("Bring anything", "PDF, Word, PowerPoint, spreadsheet, web link or pasted text."),
         ("Watch the story", "An illustrated road-trip through the ideas, drawn from your own document."),
         ("Play it", "Missions and choices with consequences. Speak or type, Urdu or English."),
         ("See progress", "Concepts unlock as understanding grows. No pass mark anywhere.")]
for i, (h, sub) in enumerate(steps):
    y = 3.4 + i * 0.83
    txt(s, M, y - 0.06, 0.5, 0.6, f"{i + 1}", size=26, color=GREY_L, font=HEAD, bold=True)
    txt(s, M + 0.58, y, 5.05, 0.8, [(h, {"size": 14.5, "bold": True, "color": NAVY, "font": HEAD, "space": 2}),
                                    (sub, {"size": 11.5, "color": GREY, "space": 0})])

h1 = picture(s, "story_pie.png", 6.75, 1.05, 5.72)
picture(s, "learn_top.png", 6.75, 1.05 + h1 + 0.2, 2.78)
picture(s, "learn_urdu_top.png", 9.69, 1.05 + h1 + 0.2, 2.78)
txt(s, 6.75, 1.05 + h1 + 0.2 + 1.83, 5.72, 0.5,
    "Above: a story scene, its chart drawn from the document’s own figures. "
    "Below: the game, and the same in Urdu, right-to-left.", size=9, color=GREY)
foot(s, 1)

# ══════════════════════════════════════════════ 2 · ARCHITECTURE
s = slide(
    "ARCHITECTURE. One small Python service on Render's free tier serves both the pages and the API, "
    "so there is no second deployment and no cross-origin set-up. A document goes in; every quote in the "
    "concept map is checked against it before anything is saved; the story is written in the background "
    "while the learner reads the map; then the tutor loop runs, streaming each reply. The learner model, "
    "meaning mastery, unlocks and the adaptation decisions, is arithmetic in our own code, not the model's "
    "memory. The admin changes settings live; they apply on the next reply.")
head(s, "ARCHITECTURE", "One service, three trust boundaries")
plain(s, 1.72,
      "It reads a document, proves every idea against it, tells the story, then holds the conversation "
      "— and keeps its own arithmetic for what you have understood.", size=15.5)

# browser
box(s, M, 3.05, 2.25, 1.75, "Learner’s browser",
    ["Story stage: SVG scenes, zooming camera", "Voice in and out, Urdu and English", "Streamed replies",
     "Language, level, limits: live"], "plain")
# service
fill(s, 4.45, 2.6, 4.7, 3.35, FAINT, BLUE_L)
txt(s, 4.57, 2.66, 4.5, 0.3, [("Rehnuma service  ·  FastAPI on one free Render instance",
                              {"size": 11.5, "bold": True, "color": DEEP, "font": HEAD})])
cells = [("Ingest", "PDF · Word · PPT · CSV · link · text"),
         ("Ground", "every quote checked against the upload"),
         ("Concept map", "ideas with their source lines"),
         ("Story", "written while the map is read"),
         ("Tutor loop", "streamed, prompt-cached"),
         ("Learner model", "mastery, unlocks, adaptation: code")]
for i, (t, d) in enumerate(cells):
    cx, cy = 4.57 + (i % 3) * 1.52, 3.08 + (i // 3) * 1.27
    box(s, cx, cy, 1.42, 1.13, t, [d], "strong" if t == "Ground" else "plain", tsize=11, lsize=8.5)
txt(s, 4.57, 5.62, 4.5, 0.28, "SQLite: the verified map (never the upload), learners, turns, settings, logs",
    size=9, color=GREY)
# claude
box(s, 10.6, 3.05, 1.88, 1.75, "Claude API", ["Opus 5: the tutor (judgement)", "Sonnet 5: the storyteller (speed)",
                                             "Offline engine if it fails"], "blue")
# admin
box(s, 4.45, 6.15, 4.7, 0.68, "Admin control room",
    ["28 settings, live, no redeploy  ·  analytics  ·  exports  ·  logs"], "plain", tsize=11, lsize=9)

arrow(s, M + 2.25, 3.92, 4.45, 3.92, ("upload +", "SSE stream"), both=True, lw=1.3, ly=0.07)
arrow(s, 9.15, 3.92, 10.6, 3.92, ("cached prompt +", "structured JSON"), both=True, lw=1.4, ly=0.07)
arrow(s, 6.8, 6.15, 6.8, 5.97, None, both=True)
txt(s, M, 5.15, 3.3, 0.9, [("Why one service", {"size": 10, "bold": True, "color": NAVY, "space": 2}),
                          ("No second deployment, no cross-origin set-up: fewer things to fail in front of a panel, "
                           "and a one-step deploy to a free tier.", {"size": 9, "color": GREY, "space": 0})])
txt(s, 10.6, 5.15, 1.88, 1.1, [("Any content", {"size": 10, "bold": True, "color": NAVY, "space": 2}),
                              ("Tested on eight kinds: biology, safety policy, code, history, Urdu, CSV, .docx, .pptx.",
                               {"size": 9, "color": GREY, "space": 0})])
foot(s, 2)

# ══════════════════════════════════════════════ 3 · SECURITY & PRIVACY
s = slide(
    "SECURITY AND PRIVACY. Two risks matter most for a bank: a tutor that invents a fact about a product, "
    "and a document written to hijack it. For the first, every quote is checked against the upload and every "
    "claim must cite one; off-source questions are refused and flagged. For the second, uploaded text and chat "
    "messages are marked as data and stripped of our own prompt delimiters, and the story can only name a layout "
    "or an icon, never supply markup. Around that: signed expiring admin sessions, no default password in "
    "production, per-client rate limits that hold behind the proxy, a daily spend cap, a strict content "
    "security policy, validated settings, minimal data and one-click erasure.")
head(s, "SECURITY, PRIVACY & RELIABILITY", "Built for a bank")
plain(s, 1.72,
      "Two risks matter most: a tutor that invents a fact about a financial product, and a document "
      "written to hijack it. Everything else is layered around those.", size=15.5)

rows = [
    ("Invented facts", "Every quote is checked against the upload before it is saved; a claim must cite a real quote; "
                       "off-source questions are refused and flagged. 100% verbatim across all eight content types tested."),
    ("Hijacking through content", "Uploaded text and chat messages are marked as data and stripped of our prompt delimiters. "
                                  "The story can only name a layout or icon: it never supplies markup."),
    ("Attacks on the server", "Web fetch refuses private networks and redirects to them; file type, size and archive limits; "
                              "strict CSP with no inline or third-party script; frame, sniff and referrer protection."),
    ("Access and abuse", "Admin cookie is signed and expires; no default password in production; per-client rate limits "
                         "that hold behind the proxy (verified live); a daily spend cap."),
    ("Bad settings", "Every setting is validated on the server. A rate limit of −5 used to break every request."),
    ("Privacy", "The upload itself is never stored, only the verified map. No personal data: a first-name label and a random "
                "ID. Retention purge, and the learner or an admin can erase a session."),
    ("Reliable & observable", "Structured logs, /health, an in-app log viewer, per-turn latency and tokens. If the AI "
                              "service fails, the offline engine takes over and says so."),
]
y0 = 2.72
for i, (k, v) in enumerate(rows):
    y = y0 + i * 0.6
    fill(s, M, y - 0.05, CW, 0.01, GREY_L)
    txt(s, M, y, 2.55, 0.5, k, size=11.5, color=NAVY, font=HEAD, bold=True)
    txt(s, M + 2.7, y, CW - 2.7, 0.55, v, size=10.5, color=GREY, spacing=1.0)
foot(s, 3)

# ══════════════════════════════════════════════ 4 · EFFICIENCY, EFFECTIVENESS & MEASUREMENT
s = slide(
    "EFFICIENCY, EFFECTIVENESS AND MEASUREMENT. Fast: replies stream, so the first words arrive in about three "
    "seconds instead of after the whole answer; the story is written while the learner reads the concept map, so it "
    "is usually ready when they press Start; about seventy percent of the prompt is served from cache; and each job "
    "gets the right-sized model, Opus 5 for the tutor's judgement, Sonnet 5 for narration. Effective: adaptation is "
    "decided in code from measured evidence, not the model's mood: difficulty, how much support, pace, and which "
    "concept comes next. There is no test, so measurement is six behaviours judged in every reply, each stored beside "
    "the learner's own words: applied it, corrected themselves, used the vocabulary, asked ahead, leaned on a hint, "
    "remembered. Before-and-after mastery per concept is the outcome measure, and the control room shows it with "
    "engagement, the learning curve, where learners struggle, helpfulness, latency, cache hits and verified quotes, "
    "all filterable and exportable. Honest limits: free storage resets on redeploy and an idle instance takes about "
    "forty seconds to wake.")
head(s, "EFFICIENCY, EFFECTIVENESS & MEASUREMENT", "Fast, adaptive, measured", size=28, width=6.0)
plain(s, 1.72, "Streaming and early preparation keep it quick. Adapting on measured evidence, and recording that "
               "evidence, is how it learns \u2014 and how we know.", width=5.85, size=14.5)

txt(s, M, 2.78, 5.85, 0.3, "FAST AND FRUGAL  \u00b7  measured on the live deployment", size=9.5, color=BLUE, bold=True)
stats = [("~3 s", "to the first words of a reply, streamed"), ("\u2264 12 s", "for the story, written while the map is read"),
         ("72%", "of prompt input served from cache"), ("+14%", "average mastery gain in the demo sessions")]
for i, (n, d) in enumerate(stats):
    x, y = M + (i % 2) * 2.95, 3.12 + (i // 2) * 0.78
    txt(s, x, y, 2.85, 0.7, [(n, {"size": 21, "bold": True, "color": NAVY, "font": HEAD, "space": 0}),
                             (d, {"size": 10, "color": GREY, "space": 0})])

txt(s, M, 4.78, 5.85, 0.3, "ADAPTS FROM MEASURED EVIDENCE, NOT THE MODEL\u2019S MOOD", size=9.5, color=BLUE, bold=True)
txt(s, M, 5.1, 5.85, 1.75, [
    ("Difficulty, support, pace and the next concept are set in code from hints, streaks and strong answers.",
     {"size": 10.5, "color": GREY, "space": 5}),
    ("Six behaviours judged every reply, each kept beside the learner\u2019s own words: applied it \u00b7 corrected "
     "themselves \u00b7 used the vocabulary \u00b7 asked ahead \u00b7 leaned on a hint \u00b7 remembered.",
     {"size": 10.5, "color": GREY, "space": 5}),
    ("Language, level or a limit like \u201cfive minutes, on a phone\u201d can change mid-session; the next reply obeys.",
     {"size": 10.5, "color": GREY, "space": 5}),
    ("Hosting is $0 on Render\u2019s free tier. Honest limits: an idle instance takes ~40 s to wake and free storage "
     "resets on redeploy.", {"size": 9.5, "color": GREY, "italic": True, "space": 0})])

picture(s, "admin_top.png", 6.95, 1.35, 5.55)
txt(s, 6.95, 1.35 + 5.55 * 800 / 1280 + 0.14, 5.55, 0.6,
    "The control room, read from three real learner sessions (English, Urdu, mixed): +14% mastery gain, 88% of 8 "
    "replies rated helpful, 72% cache hits, 35 of 35 quotes verified.", size=9, color=GREY)
txt(s, 6.95, 5.6, 5.55, 1.3, [
    ("WHAT MANAGEMENT SEES", {"size": 9.5, "bold": True, "color": BLUE, "space": 4}),
    ("Mastery gain before \u2192 after \u00b7 engagement and story completion \u00b7 learning curve \u00b7 where learners "
     "struggle \u00b7 helpfulness \u00b7 latency \u00b7 cache hits. Filter by learner, concept, signal or period; export "
     "CSV; print a one-page learner report.", {"size": 10, "color": GREY, "space": 0})])
foot(s, 4)

# ══════════════════════════════════════════════ 5 · HOW EVERY REQUIREMENT IS COVERED
s = slide(
    "HOW THE BRIEF IS COVERED. These are the twelve things the brief says a solution must demonstrate. Each row says "
    "what covers it here and which technology does the work. If the panel asks where a requirement is, point to its "
    "row and then show it in the running system.")
head(s, "THE BRIEF, POINT BY POINT", "How every requirement is covered", size=28)
txt(s, M, 1.5, CW, 0.3, "The twelve things the brief asks a solution to demonstrate: what covers each one here, and what powers it.",
    size=11.5, color=DEEP, font=HEAD)

REQ = [
    ("Any content \u2192 experience",
     "Any document becomes a verified concept map, an illustrated story, then a game \u2014 not a summary or a quiz.",
     "PyMuPDF \u00b7 Word/PowerPoint parsing \u00b7 CSV\u2192prose \u00b7 SSRF-guarded fetch \u00b7 Claude structured outputs"),
    ("Adaptive personalisation",
     "Difficulty, support, pace and next concept come from measured hints and answers; language, level and limits change live.",
     "Deterministic Python learner model \u00b7 per-learner overrides in SQLite"),
    ("Gamification & interaction",
     "Story road with a truck, missions, puzzles, boss, choices with consequences, XP, streaks, badges, unlocks.",
     "Vanilla-JS SVG scene engine \u00b7 state machine in code"),
    ("Assessment without tests",
     "Six behaviours judged each reply with the learner\u2019s own words as evidence; before/after mastery per concept.",
     "Claude structured JSON (signals + deltas) \u00b7 clamped arithmetic in Python"),
    ("Voice + accessibility",
     "Urdu/English speech in and out, code-switching, right-to-left Nastaliq, talk mode, text fallback, reduced motion, ARIA.",
     "Web Speech API \u00b7 Noto Nastaliq Urdu \u00b7 CSS/ARIA"),
    ("Accuracy & grounding",
     "Every quote is checked against the upload; claims cite a real quote; off-source questions are refused and flagged.",
     "grounding.py (normalised match + difflib repair) \u00b7 citation chips"),
    ("Highly configurable",
     "28 settings \u2014 level, language, tone, mechanics, thresholds, custom rules \u2014 change live, validated.",
     "Config service on SQLite \u00b7 admin control room \u00b7 pydantic"),
    ("Dashboards & reports",
     "Learner, engagement, mastery, usage and outcome views; filters; CSV exports; printable learner report.",
     "Chart.js (self-hosted) \u00b7 FastAPI analytics \u00b7 CSV/HTML export"),
    ("Efficient & effective",
     "First words in ~3 s; story prepared while the map is read; 72% prompt-cache hits; measurable mastery gain.",
     "SSE streaming \u00b7 prompt caching \u00b7 background threads \u00b7 Opus/Sonnet right-sizing"),
    ("Free cloud deployment",
     "One stateless service on a free tier, redeployed from GitHub on every push; scales sideways.",
     "Render free tier \u00b7 render.yaml \u00b7 GitHub \u00b7 uvicorn/FastAPI"),
    ("Cybersecurity & privacy",
     "Signed admin sessions, rate limits, strict CSP, validated inputs, SSRF guard, upload never stored, erase and retention.",
     "HMAC cookies \u00b7 slowapi \u00b7 CSP headers \u00b7 pydantic \u00b7 SQLite purge"),
    ("Reliable & observable",
     "JSON logs, /health, in-app event log, per-turn latency and tokens; offline engine takes over if the AI service fails.",
     "Python logging \u00b7 events table \u00b7 /health \u00b7 automatic fallback engine"),
]
y0, rh = 1.93, 0.4
txt(s, M, y0, 2.25, 0.2, "REQUIREMENT", size=8.5, color=BLUE, bold=True)
txt(s, M + 2.3, y0, 5.5, 0.2, "HOW IT IS COVERED HERE", size=8.5, color=BLUE, bold=True)
txt(s, M + 7.95, y0, CW - 7.95, 0.2, "TECHNOLOGY", size=8.5, color=BLUE, bold=True)
for i, (a, b, c) in enumerate(REQ):
    y = y0 + 0.24 + i * rh
    fill(s, M, y - 0.03, CW, 0.008, GREY_L)
    txt(s, M, y, 2.25, rh, a, size=10, color=NAVY, font=HEAD, bold=True)
    txt(s, M + 2.3, y, 5.5, rh, b, size=8.8, color=GREY, spacing=0.92)
    txt(s, M + 7.95, y, CW - 7.95, rh, c, size=8.5, color=DEEP, spacing=0.92)
foot(s, 5)

prs.save(OUT)
print("saved:", OUT, "| slides:", len(prs.slides))
